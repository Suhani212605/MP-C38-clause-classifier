"""Fill the college report template. Re-run after training to insert real numbers and figures:
    python fill_report.py LLM_Mini_project.docx MP-C38_Report.docx
Anything not yet produced is marked PENDING (never invented)."""
import sys, os, json, copy, glob, re
import pandas as pd
from docx import Document
from docx.shared import Pt, RGBColor, Inches

src, dst = sys.argv[1], sys.argv[2]
d = Document(src); P = list(d.paragraphs); T = d.tables
J = lambda p: json.load(open(p)) if os.path.exists(p) else {}
R, DI = J("report_assets/results.json"), J("data/data_info.json")
PEND = "PENDING (run train.py)"
g = lambda *k: (lambda x: x if x is not None else PEND)(__import__("functools").reduce(lambda a, b: a.get(b) if isinstance(a, dict) else None, k, R))
fx = lambda v: f"{v:.4f}" if isinstance(v, float) else str(v)
a_, b_ = R.get("after_test", {}), R.get("before_test", {})
N = lambda m, k: fx(m[k]) if k in m else PEND
TITLE, BLACK = "Contract Clause Type Identifier", RGBColor(0, 0, 0)


def settext(p, text, size=None):
    for r in p.runs[1:]: r._element.getparent().remove(r._element)
    r = p.runs[0] if p.runs else p.add_run(); r.text = text
    r.font.italic = False; r.font.color.rgb = BLACK
    if size: r.font.size = size


def cell(c, text): settext(c.paragraphs[0], text) if c.paragraphs[0].runs else c.paragraphs[0].add_run(text)
def image_or_pending(c, path, note):
    p = c.paragraphs[0]
    for r in p.runs: r._element.getparent().remove(r._element)
    if os.path.exists(path): p.add_run().add_picture(path, width=Inches(5.3))
    else: p.add_run(f"[PENDING: {note}]")


# ---- title page / certificate ----
settext(P[6], TITLE.upper(), Pt(18)); settext(P[7], "Project ID: MP-C38")
for i, t in enumerate(["SUHANI S RAO", "1GA23IS166", "7B"]): cell(T[0].rows[0].cells[i], t)
for r in P[28].runs:
    r.text = (r.text.replace("[Title of the Mini Project]", TITLE).replace("[MP-XXX]", "MP-C38")
              .replace("[STUDENT NAME]", "SUHANI S RAO").replace("[USN]", "1GA23IS166").replace("[7A/7B]", "7B"))
cell(T[1].rows[0].cells[0], "SDG 16: Peace, Justice and Strong Institutions    |    SDG 8: Decent Work and Economic Growth")

# ---- abstract ----
f1 = N(a_, "macro_f1")
res_s = (f"On the held-out test set the model reached an accuracy of {N(a_, 'accuracy')} and a macro-F1 of {f1}, compared with "
         f"{N(b_, 'accuracy')} and {N(b_, 'macro_f1')} before fine-tuning.") if a_ else "Final test accuracy and macro-F1 are PENDING until train.py is run."
n_cls = DI.get("per_class_used", "800")
settext(P[34], "Legal contracts contain many clauses, and manually labelling each clause (for example governing law, notices or counterparts) "
 "is slow. This project built a Contract Clause Type Identifier that predicts the type of a clause from its text. The LEDGAR dataset from the "
 f"LexGLUE benchmark was used. The 10 most frequent clause types were selected and a balanced subset of {n_cls} clauses per class was prepared "
 "(70% training, 15% validation, 15% test, random seed 42). The GPT-2 124M model from the course was re-implemented, pretrained weights were loaded, and a new "
 "10-class linear classification head was added. The head, the final layer norm and the last transformer blocks were fine-tuned with AdamW and cross-entropy loss, "
 f"and the best checkpoint was chosen by validation macro-F1. {res_s} A Gradio application classifies a single clause or every clause of an uploaded PDF/TXT contract "
 "and exports the results as CSV. The system is an academic prototype and not legal advice.")
for r in P[36].runs[1:]: r._element.getparent().remove(r._element)
P[36].add_run("GPT-2, fine-tuning, clause classification, LEDGAR, legal NLP, Gradio")

# ---- lists of figures/tables ----
settext(P[39], "")
for i in range(40, 52): P[i].runs[-1].text = P[i].runs[-1].text.replace("[  ]", "")
settext(P[40], "Figure 2.1   Class distribution of the balanced training set\t"); settext(P[42], "Figure 4.2   Confusion matrix on the test set\t")
settext(P[43], "Figure 5.1   Screenshot of the application\t")

# ---- Chapter 1 ----
settext(P[57], "Project MP-C38 – Contract Clause Type Identifier. A legal contract is a list of clauses such as governing law, notices, counterparts or severability. "
 "Lawyers, students and small businesses must find and label these clauses when they review agreements. This project builds an application where the user pastes a clause "
 "(or uploads a contract) and the fine-tuned GPT-2 model returns the predicted clause type, written as clause_type(text) → predicted type. The output is only a support tool "
 "and must be reviewed by a qualified person.")
for i, t in zip((60, 61, 62), ["Prepare a balanced LEDGAR subset of the 10 most frequent clause types (800 examples per class where available) with a stratified 70/15/15 split.",
                                "Fine-tune GPT-2 (124M) with a new 10-class head and report test accuracy and macro-F1 before and after fine-tuning, with a confusion matrix.",
                                "Provide a clause_type(text) function and a Gradio application that also classifies every clause of an uploaded PDF/TXT contract."]): settext(P[i], t)
settext(P[65], "Category MP-C: classification fine-tuning of GPT-2 with a new output head (Chapter 6 of the course). The pretrained GPT-2 small model (124M parameters) is loaded into "
 "the course-style GPTModel. The language-model head (768 → 50257) is replaced by a new Linear(768 → 10) head. Most weights are frozen; the last transformer blocks, the final "
 "layer norm and the head are trained. Macro-F1 is the principal metric because the classes must be treated equally.")

# ---- Chapter 2 ----
settext(P[69], "LEDGAR was created by Tuggener et al. (2020) from contracts filed on the US SEC EDGAR system; each example is one contract provision (text) with a "
 "clause-type label. It is included in LexGLUE (Chalkidis et al., 2022) with 100 labels and 60,000 / 10,000 / 10,000 train / validation / test examples. "
 "Two sample records: PENDING (print with: python -c \"import pandas as pd; print(pd.read_csv('data/train.csv').sample(2, random_state=1))\").")
ds = [["Dataset name and source (link)", "LEDGAR in LexGLUE (config 'ledgar'): https://huggingface.co/datasets/coastalcph/lex_glue"],
      ["Licence / terms of use", "See the dataset card (not verified in this report); used for academic work only"],
      ["Original size", "80,000 provisions, 100 labels (60k/10k/10k) per LexGLUE" + (f"; measured: {DI['original_splits']}" if DI else "")],
      ["Subset used and how it was sampled", (f"All splits pooled, duplicates removed ({DI['examples_before_dedup']} → {DI['examples_after_dedup']}), top-10 labels, "
        f"{DI['per_class_used']} random examples per class, seed {DI['seed']}") if DI else "Top-10 labels, 800 random examples per class, seed 42 (counts PENDING: run data_preparation.py)"],
      ["Labels / input-output format", "Input: clause text. Output: one of 10 clause types (class names: " + (", ".join(DI["label_map"]) if DI else "PENDING") + ")"],
      ["Train / validation / test sizes", str(DI["split_sizes"]) if DI else "5,600 / 1,200 / 1,200 if 800 per class are available (PENDING confirmation)"],
      ["Maximum token length used", "256 tokens (longer clauses truncated, shorter padded with 50256)"]]
for i, (k, v) in enumerate(ds, 1): cell(T[3].rows[i].cells[1], v)
image_or_pending(T[4].rows[0].cells[0], "report_assets/class_distribution.png", "run train.py to create class_distribution.png")
settext(P[72], "Figure 2.1: Class distribution of the balanced training set")
for i, t in zip((75, 76), ["Pool the LEDGAR train, validation and test splits; remove very short texts (< 20 characters) and collapse whitespace.",
                           "Remove duplicates (case, digits and punctuation ignored) before splitting so no clause appears in two splits; choose the 10 most frequent labels; sample the same number of clauses per class with seed 42; stratified 70/15/15 split."]): settext(P[i], t)
I = P[76]
for t in ["Counts before/after each step are saved in data/data_info.json. Only exact (normalised) duplicates are removed; near-duplicates with different wording may remain."]:
    n = copy.deepcopy(I._element); I._element.addnext(n)
    from docx.text.paragraph import Paragraph; settext(Paragraph(n, I._parent), t)
settext(P[79], "Text is converted to token IDs with the GPT-2 byte-pair-encoding tokenizer (tiktoken 'gpt2'). ClauseDataset truncates each clause to 256 tokens and pads shorter ones with token 50256 "
 "(<|endoftext|>) so a batch has one length; it also stores the real length. The DataLoader (batch size 16, shuffled with a fixed seed for training) returns token IDs, lengths and labels. "
 "The same encode/pad functions are used in inference.py so the application sees exactly what the model saw in training.")

# ---- Chapter 3 ----
settext(P[83], "The course GPTModel was re-implemented in gpt_model.py: token and positional embeddings, 12 transformer blocks (causal multi-head self-attention, feed-forward network with GELU, layer norm, residual connections), "
 "a final layer norm and an output head. The GPT-2 small (124M) pretrained weights are downloaded as a safetensors file and copied into this model; no Hugging Face model class (AutoModel, Trainer, pipeline) is used for the core model.")
cfg = [("Base model", "GPT-2 small (124M), pretrained weights from openai-community/gpt2"), ("vocab_size", "50257"), ("context_length", "1024 (clauses truncated to 256 tokens)"),
       ("emb_dim", "768"), ("n_heads", "12"), ("n_layers", "12"), ("drop_rate", "0.1"), ("qkv_bias", "True"),
       ("Output head", "Linear(768 → 10), new and randomly initialised"), ("Trainable layers", f"last {R.get('settings', {}).get('unfreeze', 2)} transformer blocks + final LayerNorm + output head"),
       ("Total / trainable parameters", f"{R['total_params']:,} / {R['trainable_params']:,}" if R else PEND)]
for i, (k, v) in enumerate(cfg, 1): cell(T[5].rows[i].cells[1], v)
settext(P[87], "The language-model head is replaced by a new Linear(768 → 10) head. All pretrained weights are first frozen; then the last transformer blocks, the final layer norm and the new head are unfrozen, "
 "so that training is fast and the general language knowledge of GPT-2 is kept. Because GPT-2 is a causal model, only the last token has seen the whole clause, so the output at the last real (non-padding) token is used as the clause representation. "
 "Padding on the right never affects earlier tokens because of the causal mask.")
S = R.get("settings", {})
ts = [("Optimizer", "AdamW (cross-entropy loss)"), ("Learning rate / weight decay", f"{S.get('lr', 5e-5)} / {S.get('wd', 0.1)}"), ("Batch size", str(S.get("batch", 16))),
      ("Number of epochs", str(S.get("epochs", 4))), ("Random seed", f"{S.get('seed', 42)} (torch, numpy, random)"),
      ("Hardware", R.get("device", "PENDING (Colab T4 GPU recommended)")), ("Training time", f"{R['train_seconds']:.0f} s" if "train_seconds" in R else PEND)]
for i, (k, v) in enumerate(ts, 1): cell(T[6].rows[i].cells[1], v)

# ---- Chapter 4 ----
for i, (k, key) in enumerate([("Test accuracy", "accuracy"), ("Macro-F1 (principal metric)", "macro_f1"), ("Test loss", "test_loss")], 1):
    cell(T[7].rows[i].cells[0], k); cell(T[7].rows[i].cells[1], N(b_, key)); cell(T[7].rows[i].cells[2], N(a_, key))
H = R.get("history", [])
settext(P[97], (f"Training loss fell from {H[0]['train_loss']:.3f} to {H[-1]['train_loss']:.3f} and validation loss went from {H[0]['val_loss']:.3f} to {H[-1]['val_loss']:.3f} over {len(H)} epochs. "
  f"Best validation macro-F1 was {R['best_val_macro_f1']:.4f}. Overfitting is indicated if validation loss rises while training loss keeps falling.") if H else "PENDING: describe the curve after running train.py.")
image_or_pending(T[8].rows[0].cells[0], "report_assets/loss_curve.png", "run train.py to create loss_curve.png")
tasks = [("(1) Balanced dataset: 10 most frequent clause types, 800 per class", "data_preparation.py with stratified 70/15/15 split", f"{DI['per_class_used']} per class, splits {DI['split_sizes']}" if DI else PEND),
         ("(2) clause_type(text) → predicted clause type", "ClauseClassifier.clause_type() in inference.py, used by the app", "Implemented; outputs PENDING until the model is trained"),
         ("(3) Report test macro-F1", "Computed with scikit-learn on test predictions", f"Macro-F1 = {f1}; macro-precision = {N(a_, 'macro_precision')}; macro-recall = {N(a_, 'macro_recall')}")]
for i, row in enumerate(tasks, 1):
    for j, v in enumerate(row): cell(T[9].rows[i].cells[j], v)
image_or_pending(T[10].rows[0].cells[0], "report_assets/confusion_matrix.png", "run train.py to create confusion_matrix.png")
settext(P[102], "Figure 4.2: Confusion matrix on the test set (per-class precision/recall/F1: report_assets/test_per_class.csv)")
err = pd.read_csv("report_assets/errors.csv").head(3) if os.path.exists("report_assets/errors.csv") else None
for i in range(1, 4):
    row = [err.iloc[i - 1].text, err.iloc[i - 1].expected, err.iloc[i - 1].predicted, "Analyse after reading the clause (ambiguous wording / similar class)"] if err is not None and len(err) >= i else ["PENDING"] * 4
    for j, v in enumerate(row): cell(T[11].rows[i].cells[j], v)

# ---- Chapter 5 ----
settext(P[109], "The Gradio application (app.py) loads the saved checkpoint models/best.pt once at start-up; it never retrains. Tab 1: the user types a clause and presses Predict to see the predicted type, the model confidence "
 "(softmax probability, not legal certainty) and the top three classes. Tab 2: the user uploads a PDF or TXT contract; the text is extracted, split into clauses (numbering, headings, paragraphs), each clause is classified, "
 "and a table plus a downloadable CSV is shown. Clear error messages are given for empty input, unsupported files and extraction failures. A legal disclaimer is displayed.")
image_or_pending(T[12].rows[0].cells[0], "report_assets/app_screenshot.png", "after running app.py, save a real screenshot as report_assets/app_screenshot.png")
SAMPLES = ["This Agreement shall be governed by the laws of the State of Delaware.", "Any notice under this Agreement shall be in writing and sent by registered mail to the address of the party.",
           "This Agreement may be executed in two or more counterparts, each of which is an original.", "If any provision is held unenforceable, the remaining provisions remain in effect.",
           "No waiver of any breach shall be deemed a waiver of any later breach."]
ap = json.load(open("report_assets/app_samples.json")) if os.path.exists("report_assets/app_samples.json") else {}
for i, s in enumerate(SAMPLES, 1):
    cell(T[13].rows[i].cells[1], s); cell(T[13].rows[i].cells[2], ap.get(s, "PENDING")); cell(T[13].rows[i].cells[3], "PENDING")

# ---- Chapter 6 ----
settext(P[117], (f"A GPT-2 (124M) based 10-class clause classifier with a new classification head was built, trained and evaluated. On the held-out test set it achieved accuracy {N(a_, 'accuracy')} and macro-F1 {f1} "
  f"(before fine-tuning: {N(b_, 'accuracy')} and {N(b_, 'macro_f1')}). A Gradio application classifies single clauses and whole contracts. Objectives 1–3 are met as shown in Table 4.2.") if a_ else
  "The complete pipeline (data preparation, GPT-2 classifier, training, evaluation, application) is implemented. Final numbers are PENDING until train.py is run; update this paragraph with the real accuracy and macro-F1.")
settext(P[120], "Only the 10 most frequent clause types are supported; other clause types will be forced into one of these 10 classes.")
settext(P[121], "Clauses are cut to 256 tokens, the model is small (124M), LEDGAR is US-English contract text, and similar classes (e.g. waiver vs. amendment) can be confused. Predictions may be wrong and are not legal advice.")
n = copy.deepcopy(P[121]._element); P[121]._element.addnext(n); from docx.text.paragraph import Paragraph
settext(Paragraph(n, P[121]._parent), "AI-tool declaration: an AI coding assistant (Claude) was used to generate the project code and draft this report; the student is responsible for understanding and explaining all code.")
settext(P[124], "Use GPT-2 medium or unfreeze more layers, and train on more LEDGAR classes with more data."); settext(P[125], "Add OCR for scanned PDFs, a better clause splitter and show the most similar training clauses as an explanation.")

# ---- references ----
refs = ['[4]  M. Tuggener, P. von Däniken, T. Peetz and M. Cieliebak, “LEDGAR: A Large-Scale Multi-label Corpus for Text Classification of Legal Provisions in Contracts,” in Proc. 12th Language Resources and Evaluation Conference (LREC), 2020, pp. 1235–1241.',
        '[5]  I. Chalkidis et al., “LexGLUE: A Benchmark Dataset for Legal Language Understanding in English,” in Proc. 60th Annual Meeting of the ACL, 2022. Available: https://huggingface.co/datasets/coastalcph/lex_glue',
        '[6]  A. Paszke et al., “PyTorch: An Imperative Style, High-Performance Deep Learning Library,” in Advances in Neural Information Processing Systems (NeurIPS), 2019.',
        '[7]  A. Abid et al., “Gradio: Hassle-Free Sharing and Testing of ML Models in the Wild,” arXiv:1906.02569, 2019.']
settext(P[131], refs[0]); settext(P[132], refs[1])
last = P[132]
for t in refs[2:]:
    n = copy.deepcopy(last._element); last._element.addnext(n); last = Paragraph(n, last._parent); settext(last, t)

# ---- appendix ----
settext(P[138], "Main file: train.py (project folder MP-C38); notebook: MP-C38_clause_classifier.ipynb")
settext(P[139], "Python 3.10+, torch>=2.1, tiktoken, datasets==2.21.0, scikit-learn, gradio (full list: requirements.txt; exact versions are printed by pip in Colab)")
for i, t in zip(range(142, 146), ["pip install -r requirements.txt", "python data_preparation.py   (downloads LEDGAR and creates data/*.csv)",
                                   "python train.py   (trains, saves models/best.pt, results and charts in report_assets/); evaluate again with: python evaluate.py",
                                   "python app.py   (opens the Gradio application; tests: pytest tests)"]): settext(P[i], t)
tmpl_head, tmpl_code = P[148]._element, P[149]._element
anchor = tmpl_head
files = ["utils.py", "gpt_model.py", "data_preparation.py", "train.py", "evaluate.py", "inference.py", "app.py", "fill_report.py", "tests/test_basic.py"]
for k, fn in enumerate(files, 1):
    h = copy.deepcopy(tmpl_head); c = copy.deepcopy(tmpl_code)
    anchor.addnext(h); h.addnext(c); anchor = c
    hp, cp = Paragraph(h, P[148]._parent), Paragraph(c, P[148]._parent)
    settext(hp, f"A.3.{k}  {fn}")
    code = open(fn, encoding="utf-8").read().rstrip("\n").split("\n")
    for r in cp.runs[1:]: r._element.getparent().remove(r._element)
    r0 = cp.runs[0]; r0.text = code[0]; r0.font.color.rgb = BLACK
    for line in code[1:]:
        r0.add_break(); r0.add_text(line)
for el in (tmpl_head, tmpl_code): el.getparent().remove(el)
for i in range(150, 153): P[i]._element.getparent().remove(P[i]._element)

# ---- remove grey guidance paragraphs ----
def grey(p): return p.runs and all((r.font.color is not None and r.font.color.type is not None and str(r.font.color.rgb) == "7F7F7F" and r.font.italic) for r in p.runs)
for p in d.paragraphs:
    if grey(p) and not p.text.startswith(("Complete code repository", "Notebook link")) and "Courier" not in str(p.runs[0].font.name): p._element.getparent().remove(p._element)
d.save(dst)
left = [p.text[:70] for p in Document(dst).paragraphs if re.search(r"\[[^\]]{2,}\]", p.text) and not p.text.startswith("[")][:0]
print("saved", dst)
